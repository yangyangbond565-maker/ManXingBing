# -*- coding: utf-8 -*-
"""生成社区慢性病大数据分析系统答辩 PPT（真实截图/真实数字）。"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "screenshots"
FIGS = ROOT / "analysis" / "output"
OUT = ROOT / "presentation" / "社区慢性病大数据分析系统-答辩.pptx"

NAVY = RGBColor(10, 28, 51)
NAVY2 = RGBColor(15, 42, 77)
INK = RGBColor(16, 35, 63)
MUTED = RGBColor(100, 116, 139)
BLUE = RGBColor(37, 99, 235)
CYAN = RGBColor(6, 182, 212)
GREEN = RGBColor(16, 185, 129)
ORANGE = RGBColor(245, 158, 11)
RED = RGBColor(239, 68, 68)
WHITE = RGBColor(255, 255, 255)
LIGHT = RGBColor(243, 247, 252)
LINE = RGBColor(214, 226, 240)
PALE_BLUE = RGBColor(234, 243, 255)
PALE_GREEN = RGBColor(229, 248, 241)
PALE_ORANGE = RGBColor(255, 246, 226)

FONT = "Microsoft YaHei"
SW, SH = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)


def bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, x, y, w, h, color, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    return s


def text(slide, x, y, w, h, value, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT,
         font=FONT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    r = p.add_run()
    r.text = value
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def bullets(slide, x, y, w, h, items, size=16, color=INK, spacing=1.2):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        p.line_spacing = spacing
        r = p.add_run()
        r.text = "·  " + item
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = color
    return box


def header(slide, kicker, title, page, dark=False):
    c1 = RGBColor(125, 211, 252) if dark else BLUE
    c2 = WHITE if dark else INK
    text(slide, 0.7, 0.38, 8.0, 0.32, kicker.upper(), 10, True, c1)
    text(slide, 0.7, 0.72, 11.6, 0.62, title, 26, True, c2)
    rect(slide, 0.7, 1.42, 0.7, 0.05, CYAN)
    text(slide, 12.15, 0.4, 0.55, 0.3, str(page).zfill(2), 10, True, c1, PP_ALIGN.RIGHT)


def footer(slide, dark=False):
    c = RGBColor(120, 160, 200) if dark else RGBColor(148, 163, 184)
    text(slide, 0.7, 7.12, 8.0, 0.25, "社区慢性病大数据分析系统 · 本科实训答辩", 9, False, c)
    text(slide, 9.8, 7.12, 2.8, 0.25, "2026.10", 9, False, c, PP_ALIGN.RIGHT)


def image_fit(slide, path, x, y, w, h):
    from PIL import Image
    im = Image.open(path)
    iw, ih = im.size
    scale = min(w / iw, h / ih)
    nw, nh = iw * scale, ih * scale
    slide.shapes.add_picture(str(path), Inches(x + (w - nw) / 2), Inches(y + (h - nh) / 2),
                             width=Inches(nw), height=Inches(nh))


def card(slide, x, y, w, h, title, value, note, color=BLUE, tint=PALE_BLUE):
    rect(slide, x, y, w, h, WHITE, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(slide, x, y, 0.08, h, color)
    text(slide, x + 0.28, y + 0.18, w - 0.4, 0.28, title, 11, False, MUTED)
    text(slide, x + 0.28, y + 0.52, w - 0.4, 0.5, value, 26, True, INK)
    text(slide, x + 0.28, y + 1.08, w - 0.4, 0.3, note, 9, False, MUTED)


def table(slide, x, y, w, h, data, col_widths=None, font_size=12, header_fill=NAVY):
    rows, cols = len(data), len(data[0])
    shp = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shp.table
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = str(val)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (WHITE if i % 2 else LIGHT)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.name = FONT
                    r.font.size = Pt(font_size)
                    r.font.bold = (i == 0)
                    r.font.color.rgb = WHITE if i == 0 else INK
    if col_widths:
        for j, cw in enumerate(col_widths):
            tbl.columns[j].width = Inches(cw)
    return tbl


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


# ---------- 1 封面 ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, NAVY)
rect(s, 0, 0, 0.18, SH, CYAN)
rect(s, 0.9, 0.8, 1.2, 0.08, CYAN)
text(s, 0.9, 1.05, 6.5, 0.4, "本科数据科学与大数据技术 · 综合实训", 14, False, RGBColor(125, 211, 252))
text(s, 0.9, 1.75, 11.6, 1.1, "社区慢性病大数据", 44, True, WHITE)
text(s, 0.9, 2.75, 11.6, 1.1, "分析系统", 44, True, WHITE)
text(s, 0.9, 4.0, 10.8, 0.5, "从 10 万患者到八模块洞察 · MySQL / Spark SQL / Spring Boot / Vue 3", 16, False, RGBColor(191, 219, 254))
rect(s, 0.9, 4.85, 5.8, 0.03, RGBColor(59, 130, 246))
text(s, 0.9, 5.15, 5.5, 0.35, "答辩人：____________", 14, False, RGBColor(203, 213, 225))
text(s, 0.9, 5.55, 5.5, 0.35, "指导教师：____________", 14, False, RGBColor(203, 213, 225))
text(s, 0.9, 5.95, 5.5, 0.35, "专业：数据科学与大数据技术", 14, False, RGBColor(203, 213, 225))
text(s, 8.6, 5.3, 3.8, 0.6, "100,000", 38, True, CYAN, PP_ALIGN.RIGHT)
text(s, 8.6, 5.95, 3.8, 0.35, "名虚构患者 · 2,065,156 条记录", 11, False, RGBColor(148, 197, 255), PP_ALIGN.RIGHT)
notes(s, "各位老师好，我是……，我的实训题目是社区慢性病大数据分析系统。项目从零搭建，完成了 10 万名虚构患者、七张业务表、八个 Spark 分析模块，以及 Spring Boot 和 Vue3 的前后端联调。")

# ---------- 2 背景 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, WHITE); header(s, "BACKGROUND", "项目背景：社区慢病管理需要“看得见”的数据", 2)
rect(s, 0.7, 1.75, 5.9, 4.9, LIGHT, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 1.0, 2.0, 5.3, 0.4, "现实问题", 17, True, BLUE)
bullets(s, 1.0, 2.55, 5.2, 3.8, [
    "患者、疾病、体检、用药、随访数据分散，难以形成完整画像",
    "传统报表只能看总量，难以回答“谁高风险、为什么、趋势如何”",
    "社区医生需要把有限的管理资源投入到最需要的人群",
    "教学场景缺少可复现、可解释、可验收的完整大数据链路"
], 14)
rect(s, 6.95, 1.75, 5.65, 2.3, NAVY, None, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 7.25, 2.0, 5.1, 0.35, "项目目标", 16, True, CYAN)
text(s, 7.25, 2.45, 5.05, 1.35, "用一套可运行的工程，把“数据生成—治理—分析—服务—可视化”串起来，并留下真实证据。", 15, False, WHITE)
rect(s, 6.95, 4.3, 5.65, 2.35, PALE_GREEN, None, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 7.25, 4.55, 5.1, 0.35, "项目定位", 16, True, GREEN)
text(s, 7.25, 5.0, 5.05, 1.3, "本科实训 / 教学演示。健康评分与趋势预测是统计模型，不是临床诊断工具。", 14, False, INK)
footer(s)
notes(s, "先说为什么做。社区慢病数据分散，传统报表不够用；我希望做一条能跑通、能解释、能验收的链路。这里也明确项目定位是教学实训，不是临床产品。")

# ---------- 3 需求核实 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, WHITE); header(s, "REQUIREMENTS", "先核实需求，再动手：原文的硬指标", 3)
cards = [
    ("患者规模", "100,000", "patient_info 患者人数，不是七表行数相加", BLUE, PALE_BLUE),
    ("业务表", "7 张", "患者 / 疾病 / 体检 / 生活习惯 / 就诊 / 用药 / 随访", CYAN, PALE_BLUE),
    ("分析模块", "8 个", "清洗、画像、疾病、风险、就诊、用药、预测、评分", GREEN, PALE_GREEN),
    ("预测算法", "线性回归", "2018-2025 训练，2026-2029 预测，输出 R²", ORANGE, PALE_ORANGE),
]
for i, c in enumerate(cards):
    card(s, 0.7 + i * 3.12, 1.75, 2.85, 1.65, *c)
text(s, 0.7, 3.7, 6.1, 0.35, "需求来源与核对", 15, True, INK)
bullets(s, 0.7, 4.12, 5.9, 2.2, [
    "飞书原文已通过浏览器实际读取，提取 60 个内容块、23 个代码块",
    "七表字段、八模块列表、10 万口径、评分权重均来自原文",
    "原文 DDL 无外键，本项目补充 6 个外键与 CHECK 约束",
    "HDFS/Hive 已实机；Sqoop 1000 行 staging 已验证"
], 12)
rect(s, 6.85, 3.65, 5.75, 2.75, LIGHT, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 7.1, 3.9, 5.2, 0.35, "八模块（原文内嵌表格）", 14, True, BLUE)
text(s, 7.1, 4.35, 5.2, 1.9, "M1 数据清洗  ·  M2 患者画像\nM3 疾病分布  ·  M4 风险关联\nM5 就诊行为  ·  M6 用药分析\nM7 趋势预测  ·  M8 健康评分", 13, False, INK)
footer(s)
notes(s, "需求不是拍脑袋。我通过浏览器读到了飞书原文，确认了 10 万是患者人数，八模块和评分权重也都有出处。原文没有外键，我补上了，但没有改字段含义。")

# ---------- 4 总体架构 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, WHITE); header(s, "ARCHITECTURE", "总体架构：一条可验证的数据链路", 4)
# boxes
layers = [
    ("数据源", "Python 生成器\n100,000 患者", 0.7, 1.8, 2.2, 1.25, BLUE, PALE_BLUE),
    ("业务库", "MySQL 8.0\n7 张表 / 206 万行", 3.25, 1.8, 2.2, 1.25, NAVY, LIGHT),
    ("数据交换", "Python ODS 导出\nTSV / HDFS / WebHDFS", 5.8, 1.8, 2.3, 1.25, CYAN, PALE_BLUE),
    ("计算", "Spark 3.3.4\n8 个分析模块", 8.45, 1.8, 2.2, 1.25, ORANGE, PALE_ORANGE),
    ("服务", "MySQL ADS\n15 张结果表", 11.0, 1.8, 1.65, 1.25, GREEN, PALE_GREEN),
]
for title, desc, x, y, w, h, c, tint in layers:
    rect(s, x, y, w, h, tint, c, MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.15, y + 0.15, w - 0.3, 0.3, title, 11, True, c)
    text(s, x + 0.15, y + 0.52, w - 0.3, 0.6, desc, 10, False, INK)
for x in [2.92, 5.47, 8.12, 10.67]:
    text(s, x, 2.25, 0.35, 0.3, "→", 18, True, MUTED, PP_ALIGN.CENTER)
rect(s, 7.95, 3.55, 4.7, 2.0, NAVY, None, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 8.2, 3.75, 4.2, 0.3, "展示层", 11, True, CYAN)
text(s, 8.2, 4.15, 4.2, 1.2, "Vue 3 管理台 + ECharts 大屏\n14 个 REST 接口读取真实 ADS\nMatplotlib 8 张静态图", 10, False, WHITE)
bullets(s, 0.7, 3.6, 6.9, 2.7, [
    "分层：ODS 原始层 → DWD 清洗层 → ADS 服务层",
    "数据质量：7 表 30+ 检查项，报告落库/落文件",
    "可替换：local / HDFS / WebHDFS 三种同步目标",
    "结果：MySQL ADS 支持分页、筛选、详情，前端不直接跑 Spark"
], 13)
footer(s)
notes(s, "架构分五层：生成、MySQL、ODS、Spark、ADS，再加前端。结果落 MySQL ADS，是为了让页面查询稳定，不用每次打开都启动 Spark。")

# ---------- 5 技术选型 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, WHITE); header(s, "STACK", "技术选型：优先保证兼容与可复现", 5)
data = [
    ["层", "技术", "版本", "选型理由"],
    ["数据生成", "Python", "3.13.9", "随机数、CSV、pymysql，易读易改"],
    ["业务库", "MySQL", "8.0.42", "事务、外键、CHECK、窗口函数"],
    ["大数据", "Spark SQL / PySpark", "3.3.4", "与原文 3.3.x 对齐，支持 SQL/JDBC"],
    ["数仓", "Hive", "3.1.3（已实机）", "ODS/DWD/ADS 分层，课程要求"],
    ["存储", "HDFS", "3.3.6（已实机）", "原始/结果分布式存储"],
    ["后端", "Spring Boot + JdbcTemplate", "2.7.18", "REST、配置简单、Java 生态"],
    ["前端", "Vue 3 + Vite + ECharts", "3.5 / 6.1", "组件化、图表强、适合大屏"],
]
table(s, 0.7, 1.75, 11.95, 4.65, data, [1.5, 3.1, 2.2, 5.15], 12)
text(s, 0.7, 6.55, 12.0, 0.3, "版本策略：先锁定组合，再逐个验证；Python 3.13 生成、Python 3.10 运行 PySpark，避免 3.13 不兼容。", 10, False, MUTED)
footer(s)
notes(s, "技术选型的原则是兼容和可复现。Spark 3.3.4 与原文一致；后端用 Spring Boot 2.7.18；前端 Vue3。Python 3.13 不能直接跑 PySpark 3.3.4，所以我单独用 3.10。")

# ---------- 6 七表关系 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, WHITE); header(s, "DATABASE", "七张业务表：一个患者，六条业务支线", 6)
rect(s, 0.8, 2.45, 2.6, 1.0, NAVY, None, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.95, 2.6, 2.3, 0.3, "patient_info", 13, True, CYAN)
text(s, 0.95, 2.95, 2.3, 0.3, "100,000 患者", 11, False, WHITE)
rights = [
    ("disease_record", "165,087", 2.1), ("health_checkup", "300,022", 3.1),
    ("lifestyle", "100,000（1:1）", 4.1), ("visit_record", "550,070", 5.1),
    ("medication_record", "549,587", 6.1), ("follow_up", "300,390", 7.1),
]
for name, count, x in rights:
    rect(s, 5.4, x, 3.4, 0.66, LIGHT, BLUE, MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 5.55, x + 0.08, 1.85, 0.25, name, 10, True, INK)
    text(s, 7.35, x + 0.08, 1.3, 0.25, count, 10, False, MUTED, PP_ALIGN.RIGHT)
    # connector
    line = s.shapes.add_connector(1, Inches(3.4), Inches(2.95), Inches(5.4), Inches(x + 0.33))
    line.line.color.rgb = RGBColor(148, 197, 255)
    line.line.width = Pt(1)
text(s, 3.55, 1.95, 1.6, 0.3, "1 : N", 10, True, BLUE, PP_ALIGN.CENTER)
bullets(s, 9.3, 2.1, 3.2, 4.0, [
    "主键 + 外键 + UNIQUE + CHECK + 索引",
    "6 个子表 ON DELETE RESTRICT，防止孤立记录",
    "lifestyle 一人一条，其余一对多",
    "字段类型按原文，未改业务口径"
], 11)
footer(s)
notes(s, "七张表以患者为主表。lifestyle 是一对一，其余是一对多。外键不是原文有的，是我为了保证数据完整性补的；删除患者会被 RESTRICT 拦住。")

# ---------- 7 数据生成与质量 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, WHITE); header(s, "DATA QUALITY", "2,065,156 条记录，质量检查 7/7 PASS", 7)
data = [
    ["表", "行数", "检查结果"],
    ["patient_info", "100,000", "PASS"],
    ["disease_record", "165,087", "PASS"],
    ["health_checkup", "300,022", "PASS"],
    ["lifestyle", "100,000", "PASS"],
    ["visit_record", "550,070", "PASS"],
    ["medication_record", "549,587", "PASS"],
    ["follow_up", "300,390", "PASS"],
]
table(s, 0.7, 1.8, 5.9, 4.25, data, [2.2, 1.7, 1.4], 11)
rect(s, 6.95, 1.8, 5.65, 2.05, PALE_GREEN, GREEN, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 7.2, 2.0, 5.1, 0.35, "检查维度", 15, True, GREEN)
text(s, 7.2, 2.45, 5.15, 1.1, "空值 · 重复 · 孤立外键 · 非法年龄/血压/血糖/BMI/心率\n非法日期 · 金额勾稽 · 10 万口径 · 主外键存在性", 12, False, INK)
rect(s, 6.95, 4.05, 5.65, 2.0, LIGHT, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 7.2, 4.25, 5.1, 0.35, "关键结果", 15, True, BLUE)
text(s, 7.2, 4.7, 5.1, 1.1, "问题总数 0；关键问题 0；\n患者 100,000；总行数 2,065,156。\n生成 seed=20261020，可复现。", 12, False, INK)
footer(s)
notes(s, "数据质量不是只看行数。我做了空值、重复、外键、日期、金额勾稽等 30 多项检查，结果是 0 问题、7 表 PASS。数据是固定 seed 生成的，可以复现。")

# ---------- 8 数据链路 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, WHITE); header(s, "PIPELINE", "数据流转：从 MySQL 到 ADS 的真实运行", 8)
steps = [
    ("① MySQL", "7 表 / 206 万行\nLOAD DATA 100.9s", BLUE, PALE_BLUE),
    ("② HDFS ODS", "7 表 / 2,065,156 行\nWebHDFS 实际同步", CYAN, PALE_BLUE),
    ("③ Spark", "8 模块 / local[4]\n读 HDFS，写 ADS", ORANGE, PALE_ORANGE),
    ("④ ADS", "MySQL 15 表\nHDFS 12 目录", GREEN, PALE_GREEN),
    ("⑤ API + UI", "14 接口 / 11 页面\n真实联调", NAVY, LIGHT),
]
for i, (title, desc, c, tint) in enumerate(steps):
    x = 0.65 + i * 2.55
    rect(s, x, 1.9, 2.2, 1.65, tint, c, MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.16, 2.08, 1.9, 0.3, title, 13, True, c)
    text(s, x + 0.16, 2.55, 1.9, 0.8, desc, 10, False, INK)
    if i < 4:
        text(s, x + 2.22, 2.55, 0.3, 0.3, "→", 18, True, MUTED, PP_ALIGN.CENTER)
rect(s, 0.7, 3.95, 11.95, 2.35, LIGHT, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.95, 4.15, 5.6, 0.35, "真实运行证据", 15, True, BLUE)
bullets(s, 0.95, 4.6, 5.6, 1.5, [
    "Spark 作业退出码 0，8 模块全部 SUCCESS",
    "ads_health_score = 100,000；ads_trend = 188 行",
    "14/14 API 通过；前端 Console 0 errors"
], 11)
text(s, 6.85, 4.15, 5.4, 0.35, "HDFS / Hive 已实机验证", 15, True, GREEN)
bullets(s, 6.85, 4.55, 5.4, 1.15, [
    "VM 192.168.20.135；HDFS 42.31 GB 可用",
    "Hive 3.1.3：7 张 ODS 表，行数与 MySQL 一致",
    "Spark 从 HDFS 读 ODS，12 张 ADS 写 HDFS + MySQL"
], 10.5, INK)
text(s, 6.85, 5.85, 5.4, 0.3, "Sqoop：1000 行 staging PASS", 10, True, GREEN)
footer(s)
notes(s, "HDFS 和 Hive 这次已经实机跑通：HDFS 测试文件校验和一致，Hive 7 张 ODS 表行数与 MySQL 一致，Spark 从 HDFS 读取并写 12 张 ADS。Sqoop 1.4.7 也已在独立目录完成 1000 行 staging 导入，HDFS/Hive 都是 1000 行；七表全量没有扩展，不能把 staging 说成全量完成。")

# ---------- 9 八模块 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, WHITE); header(s, "MODULES", "八个模块：从“看清”到“预测”", 9)
data = [
    ["模块", "关键问题", "主要输出"],
    ["M1 数据清洗", "数据能不能用？", "去重/异常过滤/清洗报告"],
    ["M2 患者画像", "患者是谁？", "性别/年龄/省份/交叉"],
    ["M3 疾病分布", "得了什么病？", "类型/严重度/共病/年度"],
    ["M4 风险因素", "和什么有关？", "吸烟/饮酒/BMI/运动/饮食"],
    ["M5 就诊行为", "怎么就医？", "类型/费用/医保/科室/年度"],
    ["M6 用药分析", "用了什么药？", "TOP/类别/有效率/联合/疾病 TOP3"],
    ["M7 趋势预测", "未来怎么变？", "年度/月度 + 线性回归 2026-2029"],
    ["M8 健康评分", "谁需要关注？", "8 维加权 0-100 / 五级 / 高危"],
]
table(s, 0.7, 1.75, 8.4, 4.9, data, [1.9, 2.7, 3.8], 10)
rect(s, 9.35, 1.75, 3.25, 4.9, NAVY, None, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 9.6, 1.95, 2.8, 0.3, "结果规模", 12, True, CYAN)
text(s, 9.6, 2.45, 2.8, 3.5, "画像 33\n疾病 10\n严重度 30\n共病 3\n风险 70\n就诊 18\n用药 108\n趋势 188\n评分 100,000", 14, False, WHITE)
footer(s)
notes(s, "八个模块回答八个问题：能不能用、患者是谁、得了什么、和什么有关、怎么就医、用了什么药、未来怎么变、谁需要关注。每个模块都有明确的输出表。")

# ---------- 10 画像与疾病 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, WHITE); header(s, "INSIGHT 01", "患者画像与疾病分布：先看清人群", 10)
image_fit(s, SHOTS / "profile.png", 0.55, 1.65, 6.0, 4.95)
image_fit(s, FIGS / "fig1_disease_distribution.png", 6.75, 1.65, 6.05, 4.95)
footer(s)
notes(s, "左边是患者画像：男女接近一半，中老年为主。右边是疾病分布：高血压 40,912 人，糖尿病 31,890 人，冠心病 23,620 人。")

# ---------- 11 风险/就诊/用药 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, WHITE); header(s, "INSIGHT 02", "风险、就诊、用药：把业务问题落到指标", 11)
image_fit(s, SHOTS / "risk.png", 0.55, 1.65, 4.0, 4.9)
image_fit(s, SHOTS / "visits.png", 4.7, 1.65, 4.05, 4.9)
image_fit(s, FIGS / "fig6_drug_top20.png", 8.85, 1.65, 3.95, 4.9)
footer(s)
notes(s, "风险因素做的是关联，不是因果。就诊看类型、费用、科室；用药看 TOP 药、类别、有效率。三张图都来自真实 ADS 表。")

# ---------- 12 趋势预测 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, WHITE); header(s, "FORECAST", "趋势预测：原文指定的一元线性回归", 12)
image_fit(s, FIGS / "fig7_trend_forecast.png", 0.55, 1.7, 7.0, 4.85)
rect(s, 7.85, 1.7, 4.75, 2.0, PALE_ORANGE, ORANGE, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 8.1, 1.9, 4.3, 0.3, "模型与结果", 13, True, ORANGE)
text(s, 8.1, 2.3, 4.3, 1.2, "y = 3902.6548 × year − 7868580.7262\nR² = 0.7361\n训练 2018-2025，预测 2026-2029", 11, False, INK)
rect(s, 7.85, 3.95, 4.75, 2.35, LIGHT, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 8.1, 4.15, 4.25, 0.3, "预测值（例）", 12, True, BLUE)
text(s, 8.1, 4.55, 4.25, 1.4, "2026：38,197.8\n2027：42,100.5\n2028：46,003.1\n2029：49,905.8", 11, False, INK)
text(s, 0.7, 6.62, 12.0, 0.3, "局限：仅 8 个年度点；R² 是拟合优度，不是临床准确率；数据为模拟数据。", 9, False, RED)
footer(s)
notes(s, "预测严格按原文的一元线性回归，没有换成 ARIMA。斜率 3902.65，R² 0.7361，预测 2026 到 2029。要强调样本只有 8 个点，R² 不是准确率。")

# ---------- 13 健康评分 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, WHITE); header(s, "SCORE", "健康评分：8 维加权，0-100 分", 13)
weights = [("血压", 20), ("血糖", 20), ("BMI", 15), ("血脂", 15), ("心率", 10), ("吸烟", 8), ("饮酒", 7), ("运动", 5)]
for i, (name, w) in enumerate(weights):
    y = 1.85 + i * 0.5
    text(s, 0.7, y, 0.8, 0.25, name, 10, False, INK)
    rect(s, 1.55, y + 0.04, 2.2, 0.14, RGBColor(226, 235, 247))
    rect(s, 1.55, y + 0.04, 2.2 * w / 20, 0.14, CYAN if w >= 15 else BLUE)
    text(s, 3.85, y - 0.02, 0.6, 0.25, str(w) + "%", 10, True, BLUE)
image_fit(s, SHOTS / "health.png", 4.65, 1.65, 5.3, 5.15)
rect(s, 10.15, 1.8, 2.5, 2.3, PALE_GREEN, GREEN, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 10.4, 2.0, 2.0, 0.3, "实测", 11, True, GREEN)
text(s, 10.4, 2.4, 2.0, 1.5, "平均 72.8\n高危 1,718\n占比 1.72%", 12, False, INK)
rect(s, 10.15, 4.35, 2.5, 2.2, LIGHT, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 10.4, 4.55, 2.0, 0.3, "分级", 11, True, BLUE)
text(s, 10.4, 4.95, 2.0, 1.4, "优秀≥85\n良好≥70\n一般≥55\n较差≥40\n危险<40", 10, False, INK)
footer(s)
notes(s, "评分权重完全按原文：血压和血糖各 20%，BMI 和血脂各 15%，心率 10%，吸烟 8%，饮酒 7%，运动 5%。平均 72.8，高危 1,718 人。")

# ---------- 14 后端 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, WHITE); header(s, "BACKEND", "Spring Boot：14 个只读接口，真实读取 ADS", 14)
data = [
    ["接口组", "代表接口", "用途"],
    ["总览", "/api/dashboard/overview", "大屏 KPI"],
    ["画像", "/api/patient-profile", "性别/年龄/省份"],
    ["疾病", "/api/disease-distribution", "类型/严重度/共病"],
    ["风险", "/api/risk-factors", "因素×疾病"],
    ["就诊", "/api/visits", "类型/费用/科室"],
    ["用药", "/api/medications", "TOP/类别/有效率"],
    ["趋势", "/api/trend", "历史+预测"],
    ["评分", "/api/health-score/overview", "等级/分箱/高危"],
    ["患者", "/api/patients/{id}", "七表完整档案"],
]
table(s, 0.7, 1.75, 7.25, 4.75, data, [1.4, 3.2, 2.5], 10)
rect(s, 8.25, 1.75, 4.35, 4.75, NAVY, None, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 8.5, 1.95, 3.9, 0.3, "统一响应", 11, True, CYAN)
text(s, 8.5, 2.35, 3.9, 0.9, "{\n  code: 0,\n  message: \"ok\",\n  data: ...\n}", 11, False, WHITE)
text(s, 8.5, 3.55, 3.9, 0.3, "验证结果", 11, True, CYAN)
text(s, 8.5, 3.95, 3.9, 1.5, "14 / 14 接口通过\nJUnit 3 / 3 通过\nHTTP 200，code=0\n无硬编码业务数据", 11, False, WHITE)
footer(s)
notes(s, "后端 14 个只读接口，统一 code/message/data。14 个 HTTP 冒烟全部通过，JUnit 3 个也通过。所有数据来自 MySQL ADS，没有硬编码。")

# ---------- 15 前端 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, WHITE); header(s, "FRONTEND", "Vue 3：管理台 + 数据大屏 + 查询详情", 15)
image_fit(s, SHOTS / "home.png", 0.55, 1.65, 6.0, 4.95)
image_fit(s, SHOTS / "patients.png", 6.75, 1.65, 6.05, 4.95)
footer(s)
notes(s, "前端是 Vue3 + Vite + ECharts + Element Plus。左边首页，右边患者查询；还有大屏、画像、疾病、风险、就诊、用药、趋势、评分、质量页面。所有图表数据来自后端。")

# ---------- 16 测试证据 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, WHITE); header(s, "EVIDENCE", "测试与证据：每个数字都能回查", 16)
cards = [
    ("Python 测试", "6 passed", "生成/外键/日期/质量/API", GREEN, PALE_GREEN),
    ("后端测试", "3 passed", "MockMvc 真实查库", BLUE, PALE_BLUE),
    ("API 冒烟", "14 / 14", "HTTP 200 / code=0", CYAN, PALE_BLUE),
    ("页面截图", "11 张", "Playwright 实际打开", ORANGE, PALE_ORANGE),
]
for i, c in enumerate(cards):
    card(s, 0.7 + i * 3.1, 1.75, 2.85, 1.55, *c)
rect(s, 0.7, 3.65, 6.1, 2.7, LIGHT, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.95, 3.85, 5.5, 0.3, "运行证据文件", 13, True, BLUE)
text(s, 0.95, 4.3, 5.5, 1.9, "reports/hdfs_hive_verification.json\nreports/spark_run.json\nreports/api_test_report.json\nlogs/hdfs_verification.log\nlogs/hive_counts.log\nlogs/spark_hdfs_run.log", 10, False, INK)
rect(s, 7.05, 3.65, 5.55, 2.7, NAVY, None, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 7.3, 3.85, 5.0, 0.3, "测试结论", 13, True, CYAN)
text(s, 7.3, 4.3, 5.0, 1.9, "核心链路通过；HDFS/Hive 已实机，七表、八模块、API、前端、图表均有真实结果。\nSqoop 1000 行 staging 已通过；七表全量未扩展。", 11, False, WHITE)
footer(s)
notes(s, "证据不是口头说。Python 6 个测试、后端 3 个、API 14 个、11 张截图，还有 HDFS/Hive/Spark 日志。HDFS/Hive 已成功；Sqoop 只验证 1000 行 staging，没有说成七表全量完成。")

# ---------- 17 总结 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, NAVY); header(s, "SUMMARY", "总结：一条真实链路，一个范围边界", 17, dark=True)
text(s, 0.7, 1.75, 5.9, 0.35, "已完成", 14, True, CYAN)
bullets(s, 0.7, 2.2, 5.8, 3.8, [
    "10 万患者、七表、206 万条记录",
    "Spark 八模块 + MySQL 15 表 + HDFS 12 目录",
    "HDFS/Hive 实机：7 ODS + 12 ADS",
    "14 个接口 + 11 个前端页面/大屏",
    "Sqoop 1000 行 staging 验证通过",
    "8 张 Matplotlib 图 + 完整文档/答辩材料"
], 12, WHITE)
text(s, 6.95, 1.75, 5.7, 0.35, "待补 / 局限", 14, True, RGBColor(252, 165, 165))
bullets(s, 6.95, 2.2, 5.6, 3.8, [
    "Sqoop 七表全量未执行（仅 1000 行 staging）",
    "预测仅 8 个年度点，未做外部验证",
    "健康评分是规则模型，非临床工具",
    "无登录权限、无并发压测"
], 12, WHITE)
rect(s, 0.7, 5.85, 12.0, 0.75, RGBColor(15, 42, 77), RGBColor(34, 211, 238), MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.95, 6.05, 11.5, 0.35, "下一步：扩到七表 Sqoop 全量 → 增加权限与增量同步 → 用真实脱敏数据做外部验证", 12, True, WHITE)
footer(s, dark=True)
notes(s, "总结：核心链路真实跑通，HDFS/Hive 也已实机，材料齐全。Sqoop 只验证了 1000 行 staging，预测和评分也不是临床模型，这些必须如实说。下一步优先扩到七表全量，再做权限和真实数据验证。")

# ---------- 18 致谢 ----------
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, NAVY)
rect(s, 0, 0, 0.18, SH, CYAN)
text(s, 1.0, 2.0, 11.0, 0.8, "感谢各位老师", 40, True, WHITE)
text(s, 1.0, 2.95, 11.0, 0.5, "请批评指正", 22, False, RGBColor(125, 211, 252))
text(s, 1.0, 4.3, 11.0, 0.35, "社区慢性病大数据分析系统 · 本科实训答辩", 12, False, RGBColor(148, 197, 255))
text(s, 1.0, 6.6, 11.0, 0.3, "PPT 中的截图、指标、测试结果均来自本项目实际运行", 9, False, RGBColor(120, 160, 200))
notes(s, "我的汇报到此结束，感谢各位老师，请批评指正。")

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print("saved", OUT, "slides", len(prs.slides))